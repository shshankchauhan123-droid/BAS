import datetime
from decimal import Decimal
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from fastapi import HTTPException, UploadFile
import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.core.database import Base
from app.user.user_model import User
from app.user.user_permission_model import UserPermission
from app.case.case_model import Case
from app.files.file_model import File
from app.bank_transactions.bank_transaction_model import BankTransaction
from app.ncrp.ncrp_model import NcrpFile, NcrpTransaction
from app.ncrp.ncrp_parser import (
    parse_ncrp_workbook,
    parse_indian_amount,
    parse_date_value,
    format_identifier,
)
from app.ncrp.ncrp_service import (
    process_ncrp_upload,
    get_case_ncrp_files_service,
    get_ncrp_file_service,
    get_ncrp_sankey_service,
    delete_ncrp_file_service,
    reconcile_ncrp_file_service,
    repair_ncrp_file_service,
    validate_ncrp_case_access,
    validate_ncrp_upload_file,
    view_ncrp_file_service,
)
from app.ncrp.ncrp_repository import get_ncrp_transactions_paginated, get_ncrp_sankey_data
from app.user.user_permission_model import UserPermission



class TestNcrpWorkbookUploadAndStorage(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(cls.engine)
        cls.SessionLocal = sessionmaker(bind=cls.engine)
        cls.temp_dir = tempfile.mkdtemp()
        import app.ncrp.ncrp_service as ncrp_svc
        ncrp_svc.NCRP_STORAGE_ROOT = Path(cls.temp_dir) / "storage" / "ncrp"

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.temp_dir, ignore_errors=True)

    def setUp(self):
        self.db = self.SessionLocal()
        # Clean tables
        self.db.query(NcrpTransaction).delete()
        self.db.query(NcrpFile).delete()
        self.db.query(BankTransaction).delete()
        self.db.query(File).delete()
        self.db.query(Case).delete()
        self.db.query(UserPermission).delete()
        self.db.query(User).delete()
        self.db.commit()

        # Create dummy user
        self.user = User(
            id=10,
            username="investigator_test",
            email="inv_test@example.com",
            password_hash="hashed_pw",
            role="user",
        )
        self.db.add(self.user)
        self.db.commit()

        # Add permissions for dummy user
        self.user_perms = UserPermission(
            user_id=self.user.id,
            can_ncrp_statement_analysis=True,
            can_bank_statement_analysis=True,
            can_upload_files=True,
            can_delete_files=True,
            can_view_reports=True,
        )
        self.db.add(self.user_perms)
        self.db.commit()
        self.db.refresh(self.user)

        # Create dummy case
        self.case = Case(
            id=20,
            case_number="CASE-NCRP-2026-001",
            case_name="Cyber Fraud Investigation 001",
            created_by=self.user.id,
            assigned_to=self.user.id,
            status="active",
        )
        self.db.add(self.case)
        self.db.commit()

    def tearDown(self):
        self.db.close()

    def _create_sample_ncrp_workbook(self, file_path: Path):
        """Create a multi-sheet Excel workbook with different headers and row structures."""
        wb = openpyxl.Workbook()

        # Sheet 1: Money transfer 2
        ws1 = wb.active
        ws1.title = "Money transfer 2"
        ws1.append([
            "Ack No", "Account No", "Transaction ID / UTR", "Txn Date",
            "Amount (INR)", "Disputed Amount", "Bank Name", "Layer", "IFSC Code", "Remarks"
        ])
        ws1.append([
            "NCRP100200", "000123456789", "UTR99887766", "15-03-2024 10:30:00",
            "94,74,000.00", "50,000.00", "State Bank of India", "Layer 1", "SBIN0001234", "Suspect account transfer"
        ])
        ws1.append([
            "NCRP100201", "000987654321", "UTR99887767", "16-03-2024",
            "1,25,500.50", "25,000.00", "HDFC Bank", "Layer 2", "HDFC0005678", "Secondary mule transfer"
        ])
        # Blank row to test safe skipping
        ws1.append([None, None, None, None, None, None, None, None, None, None])

        # Sheet 2: Transaction put on hold
        ws2 = wb.create_sheet(title="Transaction put on hold")
        # Metadata title row before actual headers
        ws2.append(["NCRP Case Export - Transaction Put on Hold", None, None, None, None])
        ws2.append([
            "Complaint ID", "Suspect Account No", "Ref No", "Date",
            "Amount", "Hold Amount", "Bank", "Action Taken by Bank", "Date of Action"
        ])
        ws2.append([
            "NCRP100200", "000123456789", "REF-HOLD-881", "17-03-2024",
            "94,74,000.00", "94,74,000.00", "State Bank of India", "Lien marked by branch", "18-03-2024"
        ])

        # Sheet 3: Cash withdraw through cheque
        ws3 = wb.create_sheet(title="Cash withdraw through cheque")
        ws3.append([
            "Acknowledgement Number", "Account / Wallet ID", "Cheque No",
            "Txn Amount", "Branch Name", "Transaction Date"
        ])
        ws3.append([
            "NCRP100200", "000987654321", "CHQ-001234",
            "50,000.00", "Connaught Place", "19-03-2024"
        ])

        wb.save(str(file_path))

    def test_indian_amount_and_identifier_parsing(self):
        """Test Indian monetary amount conversions and string identifier preservation."""
        self.assertEqual(parse_indian_amount("94,74,000.00"), Decimal("9474000.00"))
        self.assertEqual(parse_indian_amount("₹ 1,50,000.50"), Decimal("150000.50"))
        self.assertEqual(parse_indian_amount("Rs. 500"), Decimal("500.00"))
        self.assertEqual(parse_indian_amount(12500), Decimal("12500.00"))
        self.assertEqual(parse_indian_amount(None), None)
        self.assertEqual(parse_indian_amount(""), None)

        # Leading zero preservation
        self.assertEqual(format_identifier("001234567"), "001234567")
        self.assertEqual(format_identifier(123456.0), "123456")
        self.assertEqual(format_identifier(" 00987 "), "00987")
        self.assertEqual(format_identifier(None), None)

        # Date parsing
        parsed_dt = parse_date_value("15-03-2024 14:30:00")
        self.assertIsNotNone(parsed_dt)
        self.assertEqual(parsed_dt.day, 15)
        self.assertEqual(parsed_dt.month, 3)
        self.assertEqual(parsed_dt.year, 2024)

    def test_multi_sheet_dynamic_discovery_and_row_parsing(self):
        """Test that all sheets, row numbers, and headers are parsed faithfully."""
        workbook_path = Path(self.temp_dir) / "test_ncrp.xlsx"
        self._create_sample_ncrp_workbook(workbook_path)

        result = parse_ncrp_workbook(workbook_path)
        self.assertEqual(result.sheet_count, 3)
        self.assertEqual(result.total_rows, 4)
        self.assertEqual(result.imported_rows, 4)
        self.assertEqual(result.failed_rows, 0)

        # Verify sheet names
        sheet_names = [s.sheet_name for s in result.sheets_summary]
        self.assertIn("Money transfer 2", sheet_names)
        self.assertIn("Transaction put on hold", sheet_names)
        self.assertIn("Cash withdraw through cheque", sheet_names)

        # Verify row 1 in sheet 1
        row1 = result.rows[0]
        self.assertEqual(row1.sheet_name, "Money transfer 2")
        self.assertEqual(row1.row_number, 2)  # Header was row 1, data starts at row 2
        self.assertEqual(row1.acknowledgement_no, "NCRP100200")
        self.assertEqual(row1.account_or_wallet_id, "000123456789")  # Leading zeros preserved
        self.assertEqual(row1.transaction_id_utr, "UTR99887766")
        self.assertEqual(row1.transaction_amount, Decimal("9474000.00"))
        self.assertEqual(row1.disputed_amount, Decimal("50000.00"))
        self.assertEqual(row1.bank_name, "State Bank of India")
        self.assertEqual(row1.layer, "Layer 1")
        self.assertEqual(row1.raw_data.get("IFSC Code"), "SBIN0001234")
        self.assertEqual(row1.raw_data.get("Remarks"), "Suspect account transfer")

        # Verify row in sheet 2 (Transaction put on hold)
        row_hold = [r for r in result.rows if r.sheet_name == "Transaction put on hold"][0]
        self.assertEqual(row_hold.acknowledgement_no, "NCRP100200")
        self.assertEqual(row_hold.disputed_amount, Decimal("9474000.00"))
        self.assertEqual(row_hold.raw_data.get("Action Taken by Bank"), "Lien marked by branch")

        # Verify row in sheet 3 (Cash withdraw through cheque)
        row_chq = [r for r in result.rows if r.sheet_name == "Cash withdraw through cheque"][0]
        self.assertEqual(row_chq.raw_data.get("Cheque No"), "CHQ-001234")
        self.assertEqual(row_chq.transaction_amount, Decimal("50000.00"))

    def test_full_upload_and_import_service(self):
        """Test full end-to-end service: save file, insert transactions, update metadata."""
        workbook_path = Path(self.temp_dir) / "test_upload_service.xlsx"
        self._create_sample_ncrp_workbook(workbook_path)

        with open(workbook_path, "rb") as f:
            upload_file = UploadFile(filename="NCRP_Cyber_Report.xlsx", file=f)
            resp = process_ncrp_upload(
                db=self.db,
                case_id=self.case.id,
                file=upload_file,
                current_user=self.user,
            )

        self.assertEqual(resp.case_id, self.case.id)
        self.assertEqual(resp.status, "COMPLETED")
        self.assertEqual(resp.sheet_count, 3)
        self.assertEqual(resp.imported_rows, 4)
        self.assertEqual(resp.total_rows, 4)
        self.assertEqual(len(resp.sheets_summary), 3)

        # Verify database records
        ncrp_file = self.db.query(NcrpFile).filter(NcrpFile.id == resp.id).first()
        self.assertIsNotNone(ncrp_file)
        self.assertEqual(ncrp_file.original_filename, "NCRP_Cyber_Report.xlsx")
        self.assertEqual(ncrp_file.status, "COMPLETED")
        self.assertEqual(ncrp_file.imported_rows, 4)

        # Verify transaction records in ncrp_transactions
        txns = self.db.query(NcrpTransaction).filter(NcrpTransaction.ncrp_file_id == ncrp_file.id).all()
        self.assertEqual(len(txns), 4)

        # Verify strict isolation: files and bank_transactions must be completely untouched
        self.assertEqual(self.db.query(File).count(), 0)
        self.assertEqual(self.db.query(BankTransaction).count(), 0)

        # Verify listing files for case
        files = get_case_ncrp_files_service(self.db, self.case.id, self.user)
        self.assertEqual(len(files), 1)

        # Verify getting single file
        f_single = get_ncrp_file_service(self.db, resp.id, self.user)
        self.assertEqual(f_single.id, resp.id)

        # Verify delete cascade
        del_resp = delete_ncrp_file_service(self.db, resp.id, self.user)
        self.assertIn("deleted successfully", del_resp["message"])
        self.assertEqual(self.db.query(NcrpFile).count(), 0)
        self.assertEqual(self.db.query(NcrpTransaction).count(), 0)

    def test_validation_and_error_handling(self):
        """Test rejection of non-existent cases and invalid file extensions."""
        # Non-existent case
        with self.assertRaises(HTTPException) as ctx:
            validate_ncrp_case_access(self.db, 9999, self.user)
        self.assertEqual(ctx.exception.status_code, 404)

        # Invalid file extension
        invalid_file = UploadFile(filename="statement.pdf", file=None)
        with self.assertRaises(HTTPException) as ctx:
            validate_ncrp_upload_file(invalid_file)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("only accept .xlsx and .xls", ctx.exception.detail)

    def test_database_rollback_on_failure(self):
        """Test that if transaction insertion encounters an unexpected DB error, partial records roll back and file is marked FAILED."""
        from unittest.mock import patch

        workbook_path = Path(self.temp_dir) / "test_rollback.xlsx"
        self._create_sample_ncrp_workbook(workbook_path)

        with open(workbook_path, "rb") as f:
            upload_file = UploadFile(filename="NCRP_Rollback_Test.xlsx", file=f)

            # Mock bulk_save_objects to raise an exception
            with patch.object(self.db, "bulk_save_objects", side_effect=RuntimeError("Simulated DB Disk Crash")):
                with self.assertRaises(HTTPException) as ctx:
                    process_ncrp_upload(
                        db=self.db,
                        case_id=self.case.id,
                        file=upload_file,
                        current_user=self.user,
                    )
                self.assertEqual(ctx.exception.status_code, 500)

        # Verify that zero transactions were committed
        self.assertEqual(self.db.query(NcrpTransaction).count(), 0)

        # Verify file record was updated to FAILED with error details
        failed_file = self.db.query(NcrpFile).filter(NcrpFile.original_filename == "NCRP_Rollback_Test.xlsx").first()
        self.assertIsNotNone(failed_file)
        self.assertEqual(failed_file.status, "FAILED")
        self.assertIn("Simulated DB Disk Crash", failed_file.error_message)

    def test_duplicate_submission_prevention(self):
        """Test that submitting an identical file while one is in PROCESSING state is rejected with 409 Conflict."""
        # Create a file stuck in PROCESSING state
        pending_file = NcrpFile(
            case_id=self.case.id,
            created_by=self.user.id,
            original_filename="NCRP_Concurrent.xlsx",
            stored_filename="concurrent_test.xlsx",
            file_path=str(Path(self.temp_dir) / "concurrent_test.xlsx"),
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            file_size=1024,
            status="PROCESSING",
        )
        self.db.add(pending_file)
        self.db.commit()

        # Try uploading same file with same size
        dummy_file = Path(self.temp_dir) / "concurrent_test.xlsx"
        with open(dummy_file, "wb") as f:
            f.write(b"0" * 1024)

        with open(dummy_file, "rb") as f:
            upload_file = UploadFile(filename="NCRP_Concurrent.xlsx", file=f)
            with self.assertRaises(HTTPException) as ctx:
                process_ncrp_upload(
                    db=self.db,
                    case_id=self.case.id,
                    file=upload_file,
                    current_user=self.user,
                )
            self.assertEqual(ctx.exception.status_code, 409)
            self.assertIn("currently being processed", ctx.exception.detail)

    def test_alembic_migration_schema_upgrade_and_downgrade(self):
        """Test the new Alembic migration script on an isolated SQLite engine."""
        import importlib.util
        from alembic.migration import MigrationContext
        from alembic.operations import Operations

        migration_file = Path(__file__).resolve().parent.parent / "alembic" / "versions" / "f1b98a3c2e74_create_ncrp_tables.py"
        self.assertTrue(migration_file.exists(), f"Migration file {migration_file} must exist")

        spec = importlib.util.spec_from_file_location("ncrp_migration", str(migration_file))
        migration_mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(migration_mod)

        # Create fresh isolated SQLite engine
        migration_engine = create_engine("sqlite:///:memory:")

        # We first need dummy 'users' and 'cases' tables so foreign keys don't fail in SQLite
        with migration_engine.connect() as conn:
            Base.metadata.tables["users"].create(bind=conn)
            Base.metadata.tables["cases"].create(bind=conn)
            conn.commit()

            # Execute migration upgrade
            ctx = MigrationContext.configure(conn)
            op = Operations(ctx)
            # Patch alembic.op
            migration_mod.op = op
            migration_mod.upgrade()
            conn.commit()

            # Check that ncrp_file and ncrp_transactions exist
            from sqlalchemy import inspect
            inspector = inspect(conn)
            tables = inspector.get_table_names()
            self.assertIn("ncrp_file", tables)
            self.assertIn("ncrp_transactions", tables)

            # Check columns in ncrp_transactions
            cols = [c["name"] for c in inspector.get_columns("ncrp_transactions")]
            self.assertIn("sheet_name", cols)
            self.assertIn("row_number", cols)
            self.assertIn("acknowledgement_no", cols)
            self.assertIn("account_or_wallet_id", cols)
            self.assertIn("transaction_id_utr", cols)
            self.assertIn("transaction_amount", cols)
            self.assertIn("disputed_amount", cols)
            self.assertIn("bank_name", cols)
            self.assertIn("layer", cols)
            self.assertIn("raw_data", cols)

            # Execute downgrade
            migration_mod.downgrade()
            conn.commit()
            tables_after = inspect(conn).get_table_names()
            self.assertNotIn("ncrp_transactions", tables_after)
            self.assertNotIn("ncrp_file", tables_after)

    def test_ncrp_view_and_download_service(self):
        """Test retrieving original NCRP Excel workbook path and handling missing files."""
        workbook_path = Path(self.temp_dir) / "test_view.xlsx"
        self._create_sample_ncrp_workbook(workbook_path)

        with open(workbook_path, "rb") as f:
            upload_file = UploadFile(filename="NCRP_View_Test.xlsx", file=f)
            resp = process_ncrp_upload(
                db=self.db,
                case_id=self.case.id,
                file=upload_file,
                current_user=self.user,
            )

        # 1. View file with valid access
        f_path, filename, media_type = view_ncrp_file_service(
            db=self.db,
            file_id=resp.id,
            user=self.user,
            case_id=self.case.id,
        )
        self.assertTrue(f_path.exists())
        self.assertEqual(filename, "NCRP_View_Test.xlsx")
        self.assertIn("spreadsheetml", media_type)

        # 2. View file with case mismatch
        with self.assertRaises(HTTPException) as ctx:
            view_ncrp_file_service(
                db=self.db,
                file_id=resp.id,
                user=self.user,
                case_id=9999,
            )
        self.assertEqual(ctx.exception.status_code, 400)

        # 3. View non-existent file
        with self.assertRaises(HTTPException) as ctx:
            view_ncrp_file_service(
                db=self.db,
                file_id=8888,
                user=self.user,
            )
        self.assertEqual(ctx.exception.status_code, 404)

    def test_ncrp_transactions_filtering_search_pagination(self):
        """Test worksheet filtering, search across identifiers, pagination, and summaries."""
        workbook_path = Path(self.temp_dir) / "test_filter.xlsx"
        self._create_sample_ncrp_workbook(workbook_path)

        with open(workbook_path, "rb") as f:
            upload_file = UploadFile(filename="NCRP_Filter_Test.xlsx", file=f)
            resp = process_ncrp_upload(
                db=self.db,
                case_id=self.case.id,
                file=upload_file,
                current_user=self.user,
            )

        # 1. Paginated retrieval - all sheets
        res_all = get_ncrp_transactions_paginated(
            db=self.db,
            file_id=resp.id,
            case_id=self.case.id,
            skip=0,
            limit=50,
        )
        self.assertEqual(res_all["total"], 4)
        self.assertEqual(len(res_all["items"]), 4)
        self.assertIn("Money transfer 2", res_all["sheets"])
        self.assertIn("Cash withdraw through cheque", res_all["sheets"])
        self.assertIn("Transaction put on hold", res_all["sheets"])
        self.assertGreater(res_all["summary"]["total_amount"], 0)

        # 2. Sheet filter
        res_sheet = get_ncrp_transactions_paginated(
            db=self.db,
            file_id=resp.id,
            case_id=self.case.id,
            sheet_name="Money transfer 2",
        )
        self.assertEqual(res_sheet["total"], 2)
        for item in res_sheet["items"]:
            self.assertEqual(item.sheet_name, "Money transfer 2")

        # 3. Search filter by acknowledgement number or UTR
        res_search = get_ncrp_transactions_paginated(
            db=self.db,
            file_id=resp.id,
            case_id=self.case.id,
            search="NCRP100201",
        )
        self.assertEqual(res_search["total"], 1)
        self.assertEqual(res_search["items"][0].acknowledgement_no, "NCRP100201")

        # 4. Search filter by bank name
        res_bank = get_ncrp_transactions_paginated(
            db=self.db,
            file_id=resp.id,
            case_id=self.case.id,
            search="HDFC Bank",
        )
        self.assertEqual(res_bank["total"], 1)
        self.assertEqual(res_bank["items"][0].bank_name, "HDFC Bank")

        # 5. Direct bank_name and layer filters
        res_direct_bank = get_ncrp_transactions_paginated(
            db=self.db,
            file_id=resp.id,
            case_id=self.case.id,
            bank_name="HDFC",
        )
        self.assertEqual(res_direct_bank["total"], 1)
        self.assertIn("HDFC Bank", res_direct_bank["banks"])

        # 6. Min and Max amount filtering
        res_amt = get_ncrp_transactions_paginated(
            db=self.db,
            file_id=resp.id,
            case_id=self.case.id,
            min_amount=40000,
            max_amount=60000,
        )
        self.assertGreaterEqual(res_amt["total"], 1)
        for item in res_amt["items"]:
            self.assertGreaterEqual(item.transaction_amount, 40000)
            self.assertLessEqual(item.transaction_amount, 60000)

        # 7. Pagination check (limit=1)
        res_page1 = get_ncrp_transactions_paginated(
            db=self.db,
            file_id=resp.id,
            case_id=self.case.id,
            skip=0,
            limit=1,
        )
        self.assertEqual(len(res_page1["items"]), 1)
        self.assertEqual(res_page1["page"], 1)
        self.assertEqual(res_page1["total_pages"], 4)

    def test_cross_case_isolation_and_delete_permissions(self):
        """Test that cross-case deletion is rejected and can_delete_files permission is enforced."""
        workbook_path = Path(self.temp_dir) / "test_perm.xlsx"
        self._create_sample_ncrp_workbook(workbook_path)

        with open(workbook_path, "rb") as f:
            upload_file = UploadFile(filename="NCRP_Perm_Test.xlsx", file=f)
            resp = process_ncrp_upload(
                db=self.db,
                case_id=self.case.id,
                file=upload_file,
                current_user=self.user,
            )

        # 1. Reject delete with wrong case_id
        with self.assertRaises(HTTPException) as ctx:
            delete_ncrp_file_service(
                db=self.db,
                file_id=resp.id,
                user=self.user,
                case_id=9999,
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("does not belong to the specified case", ctx.exception.detail)

        # 2. Reject delete when user has can_delete_files = False
        restricted_user = User(
            id=35,
            username="restricted_user",
            email="restricted@example.com",
            password_hash="pw",
            role="user",
        )
        restricted_perms = UserPermission(
            user_id=35,
            can_delete_files=False,
            can_view_cases=True,
            can_upload_files=True,
        )
        restricted_user.permissions = restricted_perms
        self.db.add(restricted_user)
        self.db.add(restricted_perms)
        self.db.commit()

        with self.assertRaises(HTTPException) as ctx:
            delete_ncrp_file_service(
                db=self.db,
                file_id=resp.id,
                user=restricted_user,
                case_id=self.case.id,
            )
        self.assertEqual(ctx.exception.status_code, 403)
        self.assertIn("permission to delete files", ctx.exception.detail)

        # 3. Successful deletion with authorized user
        del_res = delete_ncrp_file_service(
            db=self.db,
            file_id=resp.id,
            user=self.user,
            case_id=self.case.id,
        )
        self.assertTrue(del_res["success"])
        self.assertEqual(self.db.query(NcrpFile).filter(NcrpFile.id == resp.id).count(), 0)
        self.assertEqual(self.db.query(NcrpTransaction).filter(NcrpTransaction.ncrp_file_id == resp.id).count(), 0)

    def test_money_transfer_two_utr_columns_and_bank_mapping(self):
        """Test Money Transfer to sheet with two distinct UTR columns, Bank/FIs, and Action Taken By bank."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Money Transfer to"
        ws.append([
            "S No.", "Acknowledgement No.", "Account No./ (Wallet /PG/PA) Id",
            "Transaction Id / UTR Number", "Bank/FIs", "Layer", "Account No",
            "Ifsc Code", "Transaction Date", "Transaction ID / UTR Number",
            "Transaction Amount", "Disputed Amount", "Reference No", "Remarks",
            "Action Taken By bank", "Date of Action", "pisnodal"
        ])
        ws.append([
            "1", "31206240003098", "50200094009630",
            "SBINR52024051522615907", "ICICI Bank", 2, "2427289030506856",
            "ICIC0000106", "15/05/2024 12:00:00 AM", "413615460775",
            " 2,09,187.00", " 2,09,187.00", "REF123", "Transfer remarks",
            "State Bank of India", "15/06/2024 14:08:PM", 0
        ])
        test_file = Path(self.temp_dir) / "test_dual_utr.xlsx"
        wb.save(str(test_file))

        res = parse_ncrp_workbook(test_file)
        self.assertEqual(len(res.rows), 1)
        r = res.rows[0]

        # Primary UTR must be the Transfer UTR
        self.assertEqual(r.transaction_id_utr, "413615460775")
        # raw_data must preserve both distinct UTRs
        self.assertEqual(r.raw_data.get("Transaction Id / UTR Number"), "SBINR52024051522615907")
        self.assertEqual(r.raw_data.get("Transaction ID / UTR Number"), "413615460775")
        self.assertEqual(r.raw_data.get("source_utr"), "SBINR52024051522615907")
        self.assertEqual(r.raw_data.get("transfer_utr"), "413615460775")

        # Bank/FIs must be mapped to bank_name
        self.assertEqual(r.bank_name, "ICICI Bank")
        # Action Taken By bank must be in raw_data, not bank_name
        self.assertEqual(r.raw_data.get("Action Taken By bank"), "State Bank of India")
        self.assertEqual(r.raw_data.get("action_taken_by_bank"), "State Bank of India")

        # Both accounts must remain distinct
        self.assertEqual(r.account_or_wallet_id, "50200094009630")
        self.assertEqual(r.raw_data.get("Account No"), "2427289030506856")
        self.assertEqual(r.raw_data.get("Ifsc Code"), "ICIC0000106")

        # Amount and Disputed Amount
        self.assertEqual(r.transaction_amount, Decimal("209187.00"))
        self.assertEqual(r.disputed_amount, Decimal("209187.00"))

    def test_all_six_worksheet_formats(self):
        """Test all 6 standard worksheet formats faithfully extracted without fabricating columns."""
        wb = openpyxl.Workbook()

        # 1. Money Transfer to
        ws1 = wb.active
        ws1.title = "Money Transfer to"
        ws1.append(["S No.", "Acknowledgement No.", "Account No./ (Wallet /PG/PA) Id", "Transaction Id / UTR Number", "Bank/FIs", "Layer", "Account No", "Ifsc Code", "Transaction Date", "Transaction ID / UTR Number", "Transaction Amount", "Disputed Amount"])
        ws1.append(["1", "31206240003098", "65062045083", "411561327081", "Punjab National Bank", 1, "1776002100002723", "punb0177600", "24/04/2024 12:00:00 AM", "411561327081", "10,000.00", "10,000.00"])

        # 2. Withdrawal through POS
        ws2 = wb.create_sheet(title="Withdrawal through POS")
        ws2.append(["S No.", "Acknowledgement No.", "Account No./ (Wallet /PG/PA) Id", "Transaction Id / UTR Number", "Withdrawal Date & Time", "Withdrawal Amount", "Disputed Amount", "Approval Code", "Merchant Name", "Action Taken By bank", "MID", "TID", "Layer"])
        ws2.append(["1", "31206240003098", "002289110130056384", "413616618166", "15/05/2024 00:00:AM", "1,00,000.00", "0.00", "39121688", "BABA LAKHI SHAH", "Ujjivan Small Finance Bank Ltd", "0039121688", "0039121688", 5])

        # 3. Other
        ws3 = wb.create_sheet(title="Other")
        ws3.append(["S No.", "Acknowledgement No.", "Account No./ (Wallet /PG/PA) Id", "Transaction Id / UTR Number", "Date", "Transaction Amount", "Action Taken By bank", "Layer"])
        ws3.append(["1", "31206240003098", "10176314558", "413616618306", "18/06/2024 10:08:59 PM", "1,00,000.00", "IDFC First Bank", 5])

        # 4. Transaction put on hold
        ws4 = wb.create_sheet(title="Transaction put on hold")
        ws4.append(["S No.", "Acknowledgement No.", "Account No./ (Wallet /PG/PA) Id", "Transaction Id / UTR Number", "Put on hold Date", "Put on hold Amount", "Action Taken By bank", "Layer"])
        ws4.append(["1", "31206240003098", "50100135715396", "411519719610", "20/06/2024 12:00:00 AM", 10000, "HDFC Bank", 2])

        # 5. Withdrawal through ATM
        ws5 = wb.create_sheet(title="Withdrawal through ATM")
        ws5.append(["S No.", "Acknowledgement No.", "Account No./ (Wallet /PG/PA) Id", "Transaction Id / UTR Number", "Withdrawal Date & Time", "Withdrawal Amount", "Disputed Amount", "ATM ID", "Place/Location of ATM", "Action Taken By bank", "Layer"])
        ws5.append(["1", "31206240003098", "1245100100006712", "413613453348", "15/05/2024 00:00:AM", "20,000.00", "0.00", "ATM ID :-D5388600", "JODHPUR", "Punjab National Bank", 2])

        # 6. Cash Withdrawal through Cheque
        ws6 = wb.create_sheet(title="Cash Withdrawal through Cheque")
        ws6.append(["S No.", "Acknowledgement No.", "Account No./ (Wallet /PG/PA) Id", "Transaction Id / UTR Number", "Account No", "Ifsc Code", "Cheque No", "Withdrawal Date & Time", "Withdrawal Amount", "Disputed Amount", "Branch Location", "Action Taken By bank", "Layer"])
        ws6.append(["1", "31206240003098", "20100023360118", "415013480140", "20100023360118", "BDBL0001518", 8, "29/05/2024 13:58:PM", "5,00,000.00", "0.00", "JAIPUR", "Bandhan Bank", 2])

        all_file = Path(self.temp_dir) / "test_all_six.xlsx"
        wb.save(str(all_file))

        res = parse_ncrp_workbook(all_file)
        self.assertEqual(res.sheet_count, 6)
        self.assertEqual(res.total_rows, 6)
        self.assertEqual(res.imported_rows, 6)

        # Check POS row
        pos_row = [r for r in res.rows if r.sheet_name == "Withdrawal through POS"][0]
        self.assertEqual(pos_row.transaction_id_utr, "413616618166")
        self.assertEqual(pos_row.transaction_amount, Decimal("100000.00"))
        self.assertEqual(pos_row.disputed_amount, Decimal("0.00"))
        self.assertIsNotNone(pos_row.transaction_date)
        self.assertEqual(pos_row.transaction_date.day, 15)
        self.assertEqual(pos_row.account_or_wallet_id, "002289110130056384")
        self.assertEqual(pos_row.raw_data.get("MID"), "0039121688")
        self.assertIsNone(pos_row.bank_name)
        self.assertEqual(pos_row.raw_data.get("Action Taken By bank"), "Ujjivan Small Finance Bank Ltd")

        # Check Hold row (ensure date is not confused with amount)
        hold_row = [r for r in res.rows if r.sheet_name == "Transaction put on hold"][0]
        self.assertEqual(hold_row.transaction_amount, Decimal("10000.00"))
        self.assertEqual(hold_row.disputed_amount, Decimal("10000.00"))
        self.assertIsNotNone(hold_row.transaction_date)
        self.assertEqual(hold_row.transaction_date.day, 20)
        self.assertEqual(hold_row.transaction_date.month, 6)

        # Check ATM row
        atm_row = [r for r in res.rows if r.sheet_name == "Withdrawal through ATM"][0]
        self.assertEqual(atm_row.transaction_amount, Decimal("20000.00"))
        self.assertEqual(atm_row.raw_data.get("ATM ID"), "ATM ID :-D5388600")

        # Check Cheque row
        chq_row = [r for r in res.rows if r.sheet_name == "Cash Withdrawal through Cheque"][0]
        self.assertEqual(chq_row.transaction_amount, Decimal("500000.00"))
        self.assertEqual(chq_row.raw_data.get("Cheque No"), 8)

    def test_inconsistent_am_pm_dates_and_utc_timezone(self):
        """Test robust handling of inconsistent AM/PM markers from portal exports."""
        test_dates = [
            ("15/05/2024 00:00:AM", 2024, 5, 15, 0, 0),
            ("18/06/2024 21:01:PM", 2024, 6, 18, 21, 1),
            ("29/05/2024 13:58:PM", 2024, 5, 29, 13, 58),
            ("15/02/2026 16:07:PM", 2026, 2, 15, 16, 7),
            ("18/06/2024 10:08:59 PM", 2024, 6, 18, 22, 8),
            ("20/06/2024 12:00:00 AM", 2024, 6, 20, 0, 0),
            ("11/07/2024 12:00:00 PM", 2024, 7, 11, 12, 0),
        ]
        for date_str, y, m, d, h, mi in test_dates:
            dt = parse_date_value(date_str)
            self.assertIsNotNone(dt, f"Failed to parse {date_str}")
            self.assertEqual(dt.year, y)
            self.assertEqual(dt.month, m)
            self.assertEqual(dt.day, d)
            self.assertEqual(dt.hour, h)
            self.assertEqual(dt.minute, mi)
            self.assertEqual(dt.tzinfo, datetime.timezone.utc)

    def test_reconcile_and_repair_workflow(self):
        """Test dry-run reconciliation and atomic repair of NCRP file."""
        # 1. First upload a file
        workbook_path = Path(self.temp_dir) / "test_repair_wf.xlsx"
        self._create_sample_ncrp_workbook(workbook_path)

        with open(workbook_path, "rb") as f:
            upload_file = UploadFile(filename="Repair_Test.xlsx", file=f)
            resp = process_ncrp_upload(
                db=self.db,
                case_id=self.case.id,
                file=upload_file,
                current_user=self.user,
            )

        # 2. Dry-run reconciliation
        rec_res = reconcile_ncrp_file_service(
            db=self.db,
            file_id=resp.id,
            user=self.user,
            case_id=self.case.id,
        )
        self.assertTrue(rec_res["dry_run"])
        self.assertFalse(rec_res["repaired"])
        self.assertEqual(rec_res["total_source_rows"], 4)
        self.assertEqual(rec_res["total_db_rows"], 4)

        # 3. Safe repair execution
        repair_res = repair_ncrp_file_service(
            db=self.db,
            file_id=resp.id,
            user=self.user,
            case_id=self.case.id,
        )
        self.assertTrue(repair_res["repaired"])
        self.assertEqual(repair_res["total_db_rows"], 4)
        self.assertEqual(repair_res["discrepancies_count"], 0)

        # Verify DB records intact after repair
        txns = self.db.query(NcrpTransaction).filter(NcrpTransaction.ncrp_file_id == resp.id).all()
        self.assertEqual(len(txns), 4)

    def test_ncrp_sankey_data_and_kpis(self):
        """Test layer-wise Sankey nodes, links aggregation, and financial KPIs."""
        workbook_path = Path(self.temp_dir) / "test_sankey_wf.xlsx"
        self._create_sample_ncrp_workbook(workbook_path)

        with open(workbook_path, "rb") as f:
            upload_file = UploadFile(filename="Sankey_Test.xlsx", file=f)
            resp = process_ncrp_upload(
                db=self.db,
                case_id=self.case.id,
                file=upload_file,
                current_user=self.user,
            )

        # Call Sankey service for all layers
        sankey_res = get_ncrp_sankey_service(
            db=self.db,
            file_id=resp.id,
            user=self.user,
            case_id=self.case.id,
        )

        self.assertIn("nodes", sankey_res)
        self.assertIn("links", sankey_res)
        self.assertIn("kpis", sankey_res)
        self.assertIn("available_layers", sankey_res)
        self.assertIn("available_acknowledgements", sankey_res)

        kpis = sankey_res["kpis"]
        # In _create_sample_ncrp_workbook there are 4 rows: 2 Money Transfer (L1, L2), 1 Hold, 1 Cheque
        self.assertEqual(kpis["total_transactions"], 4)
        self.assertGreaterEqual(kpis["unique_source_accounts"], 1)
        self.assertGreaterEqual(kpis["unique_destination_accounts"], 1)
        self.assertGreater(kpis["total_transaction_amount"], 0)
        # The Transaction put on hold row is correctly counted in unknown_or_hold_flow_count
        self.assertEqual(kpis["unknown_or_hold_flow_count"], 1)

        # Check nodes have layer prefixes and no empty accounts
        nodes = sankey_res["nodes"]
        self.assertGreater(len(nodes), 0)
        for n in nodes:
            self.assertTrue(n["id"].startswith("L"))
            self.assertIsNotNone(n["account"])
            self.assertIsNotNone(n["bank"])

        # Check links have non-zero values and source/target mapping
        links = sankey_res["links"]
        self.assertGreater(len(links), 0)
        for l in links:
            self.assertGreater(l["value"], 0)
            self.assertTrue(l["source"].startswith("L"))
            self.assertTrue(l["target"].startswith("L"))

    def test_ncrp_sankey_and_paginated_filtering(self):
        """Test layer and acknowledgement filtering across Sankey and paginated transactions."""
        workbook_path = Path(self.temp_dir) / "test_sankey_filt.xlsx"
        self._create_sample_ncrp_workbook(workbook_path)

        with open(workbook_path, "rb") as f:
            upload_file = UploadFile(filename="Sankey_Filt.xlsx", file=f)
            resp = process_ncrp_upload(
                db=self.db,
                case_id=self.case.id,
                file=upload_file,
                current_user=self.user,
            )

        # Test Layer 1 filtering on Sankey
        l1_sankey = get_ncrp_sankey_service(
            db=self.db,
            file_id=resp.id,
            user=self.user,
            case_id=self.case.id,
            layer="1",
        )
        self.assertEqual(l1_sankey["kpis"]["total_transactions"], 1)
        self.assertEqual(len(l1_sankey["links"]), 1)

        # Test Layer 2 filtering on Sankey (Money transfer row 2)
        l2_sankey = get_ncrp_sankey_service(
            db=self.db,
            file_id=resp.id,
            user=self.user,
            case_id=self.case.id,
            layer="2",
        )
        self.assertEqual(l2_sankey["kpis"]["total_transactions"], 1)

        # Test paginated transactions with acknowledgement_no filter (NCRP100200 has 3 rows)
        paginated_res = get_ncrp_transactions_paginated(
            db=self.db,
            file_id=resp.id,
            case_id=self.case.id,
            acknowledgement_no="NCRP100200",
        )
        self.assertEqual(paginated_res["total"], 3)
        self.assertIn("acknowledgements", paginated_res)
        self.assertIn("NCRP100200", paginated_res["acknowledgements"])

    def test_ncrp_money_flow_sankey_strict_money_transfer(self):
        """Test Money Flow Graph strictly restricts data to Money Transfer worksheet only and excludes non-transfer events."""
        workbook_path = Path(self.temp_dir) / "test_money_flow.xlsx"
        self._create_sample_ncrp_workbook(workbook_path)

        with open(workbook_path, "rb") as f:
            upload_file = UploadFile(filename="Money_Flow_Test.xlsx", file=f)
            resp = process_ncrp_upload(
                db=self.db,
                case_id=self.case.id,
                file=upload_file,
                current_user=self.user,
            )

        # Money Flow Graph with money_transfer_only=True
        flow_res = get_ncrp_sankey_service(
            db=self.db,
            file_id=resp.id,
            user=self.user,
            case_id=self.case.id,
            money_transfer_only=True,
        )

        kpis = flow_res["kpis"]
        # Only the 2 Money Transfer rows are eligible, Hold and Cheque are excluded
        self.assertEqual(kpis["total_transactions"], 2)
        self.assertEqual(kpis["unknown_or_hold_flow_count"], 0)
        self.assertEqual(kpis["excluded_incomplete_count"], 0)
        self.assertEqual(len(flow_res["links"]), 2)
        self.assertIn("layer_summaries", flow_res)
        self.assertGreaterEqual(len(flow_res["layer_summaries"]), 2)
        l1_sum = flow_res["layer_summaries"][0]
        self.assertEqual(l1_sum["layer"], "Layer 1")
        self.assertEqual(l1_sum["layer_num"], 1)
        self.assertGreater(l1_sum["transaction_count"], 0)
        self.assertGreater(l1_sum["total_amount"], 0)
        self.assertGreater(l1_sum["unique_senders"], 0)
        self.assertGreater(l1_sum["unique_receivers"], 0)


if __name__ == "__main__":
    unittest.main()

