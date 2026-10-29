import datetime
from decimal import Decimal
import os
from pathlib import Path
import re
from typing import Any, Dict, List, Optional, Tuple

import dateutil.parser
import openpyxl

try:
    import xlrd
except ImportError:
    xlrd = None


# ============================================================
# NORMALIZATION HELPERS
# ============================================================

def normalize_header(header: Any) -> str:
    """Normalize header text for comparison by removing punctuation and lowercasing."""
    if header is None:
        return ""
    text = str(header).strip().lower()
    return re.sub(r"[^a-z0-9]", "", text)


def parse_indian_amount(value: Any) -> Optional[Decimal]:
    """Parse monetary amounts including Indian-formatted values (e.g. 94,74,000.00, Rs. 500, ₹ 1,50,000.50).

    Preserves exact decimal precision and quantizes to 2 decimal places.
    """
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"))
        except Exception:
            return None

    str_val = str(value).strip()
    if not str_val:
        return None

    # Remove trailing "/-" (common in India: "500/-")
    str_val = re.sub(r"/-\s*$", "", str_val)
    # Remove currency prefixes like "Rs.", "Rs", "INR", "₹"
    str_val = re.sub(r"(?i)(rs\.?|inr|re\.?|₹)\s*", "", str_val)
    # Remove commas and spaces
    str_val = str_val.replace(",", "").strip()
    # Now only digits, minus, and dot should remain
    cleaned = re.sub(r"[^\d.-]", "", str_val)
    if not cleaned or cleaned in {"-", ".", "-."}:
        return None

    try:
        return Decimal(cleaned).quantize(Decimal("0.01"))
    except Exception:
        return None


def parse_date_value(value: Any) -> Optional[datetime.datetime]:
    """Parse Excel date objects, serial numbers, or date strings into UTC datetime.

    Robustly handles Indian DD/MM/YYYY date formats and portal export AM/PM inconsistencies
    such as '15/05/2024 00:00:AM', '18/06/2024 21:01:PM', '29/05/2024 13:58:PM', and '16:07:PM'.
    """
    if value is None:
        return None

    if isinstance(value, datetime.datetime):
        if value.tzinfo is None:
            return value.replace(tzinfo=datetime.timezone.utc)
        return value.astimezone(datetime.timezone.utc)

    if isinstance(value, datetime.date):
        return datetime.datetime.combine(
            value, datetime.time.min, tzinfo=datetime.timezone.utc
        )

    # Check for Excel serial number (float or int between 30000 and 70000, roughly 1982 to 2091)
    if isinstance(value, (int, float)) and 30000 <= value <= 70000:
        try:
            base_date = datetime.datetime(1899, 12, 30, tzinfo=datetime.timezone.utc)
            return base_date + datetime.timedelta(days=float(value))
        except Exception:
            pass

    s = str(value).strip()
    if not s:
        return None

    # Regex matching DD/MM/YYYY or YYYY-MM-DD or DD-MM-YYYY with optional time and AM/PM (including colon-prefixed :PM / :AM)
    m = re.match(
        r"^(\d{1,4})[/-](\d{1,2})[/-](\d{1,4})(?:[\sT]+(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?)?(?::?\s*([a-zA-Z]{2}))?$",
        s,
    )
    if m:
        p1, p2, p3, hour, minute, second, meridian = m.groups()
        try:
            if len(p1) == 4:
                year, month, day = int(p1), int(p2), int(p3)
            else:
                day, month, year = int(p1), int(p2), int(p3)

            h = int(hour) if hour is not None else 0
            mi = int(minute) if minute is not None else 0
            sec = int(second) if second is not None else 0

            if meridian:
                mer = meridian.upper()
                if mer == "PM":
                    # If 1 <= h < 12, convert to afternoon/evening 24h
                    if h < 12:
                        h += 12
                    # If h >= 12 (e.g. 16:07:PM, 21:01:PM), it is ALREADY 24-hour; keep h as is.
                elif mer == "AM":
                    # 12:xx AM is midnight hour 0
                    if h == 12:
                        h = 0
                    # If h == 0 (e.g. 00:00:AM) or 1 <= h < 12, keep h.

            return datetime.datetime(year, month, day, h, mi, sec, tzinfo=datetime.timezone.utc)
        except Exception:
            pass

    # Fallback to dateutil with colon-separated :AM / :PM normalized to space
    try:
        cleaned_s = re.sub(r":\s*([AP]M)", r" \1", s, flags=re.IGNORECASE)
        dt = dateutil.parser.parse(cleaned_s, dayfirst=True)
        if dt.tzinfo is None:
            return dt.replace(tzinfo=datetime.timezone.utc)
        return dt.astimezone(datetime.timezone.utc)
    except Exception:
        return None


def format_identifier(value: Any) -> Optional[str]:
    """Preserve identifiers (account numbers, ack nos, UTRs, IFSC, MID, TID) as strings without losing leading zeros."""
    if value is None:
        return None
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        s = f"{value:.0f}" if value.is_integer() else str(value).strip()
        return s
    if isinstance(value, int):
        return str(value)
    str_val = str(value).strip()
    return str_val if str_val else None


def serialize_for_json(val: Any) -> Any:
    """Safely serialize date, decimal, and complex objects for JSONB storage."""
    if val is None:
        return None
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.isoformat()
    if isinstance(val, Decimal):
        return str(val)
    if isinstance(val, (int, float, bool, str)):
        return val
    return str(val)


# ============================================================
# WORKSHEET PARSER DATA STRUCTURES
# ============================================================

class ParsedRow:
    def __init__(
        self,
        sheet_name: str,
        row_number: int,
        acknowledgement_no: Optional[str] = None,
        account_or_wallet_id: Optional[str] = None,
        transaction_id_utr: Optional[str] = None,
        transaction_date: Optional[datetime.datetime] = None,
        transaction_amount: Optional[Decimal] = None,
        disputed_amount: Optional[Decimal] = None,
        bank_name: Optional[str] = None,
        layer: Optional[str] = None,
        raw_data: Optional[Dict[str, Any]] = None,
    ):
        self.sheet_name = sheet_name
        self.row_number = row_number
        self.acknowledgement_no = acknowledgement_no
        self.account_or_wallet_id = account_or_wallet_id
        self.transaction_id_utr = transaction_id_utr
        self.transaction_date = transaction_date
        self.transaction_amount = transaction_amount
        self.disputed_amount = disputed_amount
        self.bank_name = bank_name
        self.layer = layer
        self.raw_data = raw_data or {}


class SheetSummary:
    def __init__(self, sheet_name: str, total_rows: int, imported_rows: int, failed_rows: int):
        self.sheet_name = sheet_name
        self.total_rows = total_rows
        self.imported_rows = imported_rows
        self.failed_rows = failed_rows

    def to_dict(self) -> Dict[str, Any]:
        return {
            "sheet_name": self.sheet_name,
            "total_rows": self.total_rows,
            "imported_rows": self.imported_rows,
            "failed_rows": self.failed_rows,
        }


class NcrpParseResult:
    def __init__(
        self,
        sheet_count: int,
        total_rows: int,
        imported_rows: int,
        failed_rows: int,
        rows: List[ParsedRow],
        sheets_summary: List[SheetSummary],
    ):
        self.sheet_count = sheet_count
        self.total_rows = total_rows
        self.imported_rows = imported_rows
        self.failed_rows = failed_rows
        self.rows = rows
        self.sheets_summary = sheets_summary


# ============================================================
# HEADER-DRIVEN WORKSHEET PARSER
# ============================================================

def parse_row_by_worksheet_rules(
    sheet_name: str,
    row_idx: int,
    raw_headers: List[str],
    row_values: List[Any],
) -> ParsedRow:
    """Parse a single data row with worksheet-aware, header-driven mapping rules."""
    norm_sheet = normalize_header(sheet_name)
    headers_norm = [normalize_header(h) for h in raw_headers]
    row_len = len(row_values)

    def get_cell_by_norm(target_norms: List[str]) -> Tuple[Optional[Any], Optional[int]]:
        for norm_target in target_norms:
            for idx, h_norm in enumerate(headers_norm):
                if h_norm == norm_target and idx < row_len:
                    return row_values[idx], idx
        return None, None

    # Base raw_data dictionary retaining every single source column with exact header name
    raw_data: Dict[str, Any] = {}
    for col_idx, cell_val in enumerate(row_values):
        if col_idx < len(raw_headers):
            h_name = raw_headers[col_idx]
            raw_data[h_name] = serialize_for_json(cell_val)

    # 1. Acknowledgement No
    ack_val, _ = get_cell_by_norm([
        "acknowledgementno", "acknowledgementno.", "ackno", "acknum",
        "complaintid", "incidentno", "tokenno", "acknowledgementnumber"
    ])
    ack_no = format_identifier(ack_val)

    # 2. Account or Wallet ID
    acc_val, _ = get_cell_by_norm([
        "accountnowalletpgpaid", "accountnowalletid", "accountorwalletid",
        "accountwalletid", "accountnowallet", "walletid", "suspectaccountno",
        "suspectaccount", "beneficiaryaccountno", "beneficiaryaccount",
        "accountno", "accountnumber", "accno"
    ])
    acc_or_wallet = format_identifier(acc_val)

    # 3. Layer
    layer_val, _ = get_cell_by_norm(["layer", "layering", "level", "hop"])
    layer_str = format_identifier(layer_val)

    # Detect worksheet classification
    is_money_transfer = "moneytransfer" in norm_sheet
    is_pos = "pos" in norm_sheet
    is_atm = "atm" in norm_sheet
    is_cheque = "cheque" in norm_sheet or "check" in norm_sheet
    is_hold = "hold" in norm_sheet
    is_other = "other" in norm_sheet

    # 4. UTR detection (distinguish multiple UTR columns in Money Transfer sheets)
    utr_col_indices = [
        idx for idx, h_norm in enumerate(headers_norm)
        if any(k in h_norm for k in ["utr", "txnid", "transactionid", "referenceno", "refno"])
        and not any(k in h_norm for k in ["date", "amount", "fee", "hold"])
    ]

    source_utr = None
    transfer_utr = None
    txn_utr = None

    if is_money_transfer and len(utr_col_indices) >= 2:
        # Col 1 is Source UTR ('Transaction Id / UTR Number')
        # Col 2 is Transfer UTR ('Transaction ID / UTR Number')
        src_val = row_values[utr_col_indices[0]] if utr_col_indices[0] < row_len else None
        tgt_val = row_values[utr_col_indices[1]] if utr_col_indices[1] < row_len else None
        source_utr = format_identifier(src_val)
        transfer_utr = format_identifier(tgt_val)
        txn_utr = transfer_utr or source_utr
    elif len(utr_col_indices) >= 1:
        u_val = row_values[utr_col_indices[0]] if utr_col_indices[0] < row_len else None
        txn_utr = format_identifier(u_val)
        if is_money_transfer:
            transfer_utr = txn_utr

    # 5. Dates
    txn_date = None
    if is_money_transfer:
        dt_val, _ = get_cell_by_norm(["transactiondate", "txndate", "dateoftransaction", "entrydate", "date"])
        txn_date = parse_date_value(dt_val)
    elif is_pos or is_atm or is_cheque:
        dt_val, _ = get_cell_by_norm(["withdrawaldatetime", "withdrawaldate", "transactiondate", "txndate", "date"])
        txn_date = parse_date_value(dt_val)
    elif is_hold:
        dt_val, _ = get_cell_by_norm(["putonholddate", "holddate", "date", "transactiondate"])
        txn_date = parse_date_value(dt_val)
    elif is_other:
        dt_val, _ = get_cell_by_norm(["date", "transactiondate", "txndate"])
        txn_date = parse_date_value(dt_val)
    else:
        # Generic date fallback
        dt_val = None
        for idx, h_norm in enumerate(headers_norm):
            if "date" in h_norm and not any(k in h_norm for k in ["action", "complaint", "birth"]):
                if idx < row_len:
                    dt_val = row_values[idx]
                    break
        txn_date = parse_date_value(dt_val)

    # 6. Amounts
    txn_amount = None
    disp_amount = None

    if is_money_transfer:
        amt_val, _ = get_cell_by_norm(["transactionamount", "txnamount", "transferamount", "amountinr", "amountrs", "inramount", "amount"])
        txn_amount = parse_indian_amount(amt_val)

        disp_val, _ = get_cell_by_norm(["disputedamount", "disputed"])
        disp_amount = parse_indian_amount(disp_val)

    elif is_pos or is_atm or is_cheque:
        amt_val, _ = get_cell_by_norm(["withdrawalamount", "txnamount", "transactionamount", "amountinr", "amount"])
        txn_amount = parse_indian_amount(amt_val)

        disp_val, _ = get_cell_by_norm(["disputedamount", "disputed"])
        disp_amount = parse_indian_amount(disp_val)

    elif is_hold:
        # In hold sheets, the hold amount is both the transaction amount and disputed/held amount
        amt_val, _ = get_cell_by_norm(["putonholdamount", "holdamount", "lienamount", "blockedamount", "amount", "transactionamount"])
        txn_amount = parse_indian_amount(amt_val)
        disp_amount = parse_indian_amount(amt_val)

    elif is_other:
        amt_val, _ = get_cell_by_norm(["transactionamount", "txnamount", "amount", "amountinr", "inramount"])
        txn_amount = parse_indian_amount(amt_val)

        disp_val, _ = get_cell_by_norm(["disputedamount", "disputed"])
        disp_amount = parse_indian_amount(disp_val)

    else:
        amt_val = None
        for idx, h_norm in enumerate(headers_norm):
            if any(k in h_norm for k in ["amount", "inr"]) and not any(k in h_norm for k in ["disputed", "hold", "lien", "blocked", "date"]):
                if idx < row_len:
                    amt_val = row_values[idx]
                    break
        txn_amount = parse_indian_amount(amt_val)

        disp_val = None
        for idx, h_norm in enumerate(headers_norm):
            if any(k in h_norm for k in ["disputed", "holdamount", "putonholdamount", "lienamount"]):
                if idx < row_len:
                    disp_val = row_values[idx]
                    break
        disp_amount = parse_indian_amount(disp_val)

    # 7. Bank Name vs Action Taken By bank
    bank_name = None
    action_bank = None

    # Check for Action Taken By bank first
    action_val, _ = get_cell_by_norm(["actiontakenbybank", "actionbank"])
    action_bank = str(action_val).strip() if action_val is not None and str(action_val).strip() else None

    # Target Bank (Bank/FIs)
    b_val, _ = get_cell_by_norm(["bankfis", "bankfi", "bankname", "beneficiarybank", "suspectbank"])
    if b_val is not None and str(b_val).strip():
        bank_name = str(b_val).strip()
    elif is_money_transfer or not action_bank:
        # In Money Transfer sheets or sheets without separate action bank, check for 'bank'
        b_generic, _ = get_cell_by_norm(["bank"])
        if b_generic is not None and str(b_generic).strip():
            bank_name = str(b_generic).strip()

    # Preserve standardized convenience fields in raw_data for full fidelity
    if source_utr is not None:
        raw_data["source_utr"] = source_utr
    if transfer_utr is not None:
        raw_data["transfer_utr"] = transfer_utr
    if bank_name is not None:
        raw_data["bank_name"] = bank_name
    if action_bank is not None:
        raw_data["action_taken_by_bank"] = action_bank

    return ParsedRow(
        sheet_name=sheet_name,
        row_number=row_idx,
        acknowledgement_no=ack_no,
        account_or_wallet_id=acc_or_wallet,
        transaction_id_utr=txn_utr,
        transaction_date=txn_date,
        transaction_amount=txn_amount,
        disputed_amount=disp_amount,
        bank_name=bank_name,
        layer=layer_str,
        raw_data=raw_data,
    )


# ============================================================
# OPENPYXL & XLRD SHEET PROCESSORS
# ============================================================

def _parse_openpyxl_sheet(sheet, sheet_name: str) -> Tuple[List[ParsedRow], int, int, int]:
    """Parse a single openpyxl worksheet."""
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], 0, 0, 0

    header_idx = -1
    raw_headers = []

    for idx, row in enumerate(rows):
        non_empty = [c for c in row if c is not None and str(c).strip()]
        if len(non_empty) >= 2:
            string_count = sum(1 for c in non_empty if not isinstance(c, (int, float)))
            if string_count >= 1:
                header_idx = idx
                raw_headers = [str(c).strip() if c is not None else f"Column_{c_idx+1}" for c_idx, c in enumerate(row)]
                break

    if header_idx == -1 or not raw_headers:
        header_idx = 0
        raw_headers = [str(c).strip() if c is not None else f"Column_{c_idx+1}" for c_idx, c in enumerate(rows[0])]

    parsed_rows: List[ParsedRow] = []
    total_eligible = 0
    imported = 0
    failed = 0

    for row_offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(c is not None and str(c).strip() for c in row):
            continue

        total_eligible += 1
        try:
            parsed = parse_row_by_worksheet_rules(
                sheet_name=sheet_name,
                row_idx=row_offset,
                raw_headers=raw_headers,
                row_values=list(row),
            )
            parsed_rows.append(parsed)
            imported += 1
        except Exception:
            failed += 1

    return parsed_rows, total_eligible, imported, failed


def _parse_xls_workbook(file_path: Path) -> NcrpParseResult:
    """Parse legacy .xls workbook using xlrd."""
    if xlrd is None:
        raise ValueError("xlrd is required to parse .xls files. Please upload .xlsx instead.")

    book = xlrd.open_workbook(str(file_path))
    all_rows: List[ParsedRow] = []
    sheets_summary: List[SheetSummary] = []
    total_discovered_rows = 0
    total_imported_rows = 0
    total_failed_rows = 0

    for sheet_name in book.sheet_names():
        sheet = book.sheet_by_name(sheet_name)
        if sheet.nrows == 0:
            sheets_summary.append(SheetSummary(sheet_name, 0, 0, 0))
            continue

        header_idx = -1
        raw_headers = []
        for r_idx in range(sheet.nrows):
            row_vals = [sheet.cell_value(r_idx, c_idx) for c_idx in range(sheet.ncols)]
            non_empty = [c for c in row_vals if c != ""]
            if len(non_empty) >= 2:
                header_idx = r_idx
                raw_headers = [str(c).strip() if c != "" else f"Column_{c_idx+1}" for c_idx, c in enumerate(row_vals)]
                break

        if header_idx == -1:
            header_idx = 0
            raw_headers = [str(sheet.cell_value(0, c)).strip() or f"Column_{c+1}" for c in range(sheet.ncols)]

        sheet_eligible = 0
        sheet_imported = 0
        sheet_failed = 0

        for r_idx in range(header_idx + 1, sheet.nrows):
            row_vals = [sheet.cell_value(r_idx, c_idx) for c_idx in range(sheet.ncols)]
            if not any(c != "" for c in row_vals):
                continue

            sheet_eligible += 1
            try:
                parsed = parse_row_by_worksheet_rules(
                    sheet_name=sheet_name,
                    row_idx=r_idx + 1,
                    raw_headers=raw_headers,
                    row_values=row_vals,
                )
                all_rows.append(parsed)
                sheet_imported += 1
            except Exception:
                sheet_failed += 1

        sheets_summary.append(SheetSummary(sheet_name, sheet_eligible, sheet_imported, sheet_failed))
        total_discovered_rows += sheet_eligible
        total_imported_rows += sheet_imported
        total_failed_rows += sheet_failed

    return NcrpParseResult(
        sheet_count=len(sheets_summary),
        total_rows=total_discovered_rows,
        imported_rows=total_imported_rows,
        failed_rows=total_failed_rows,
        rows=all_rows,
        sheets_summary=sheets_summary,
    )


def parse_ncrp_workbook(file_path: Path) -> NcrpParseResult:
    """Parse an NCRP workbook (.xlsx or .xls) across all worksheets dynamically."""
    if not file_path.exists():
        raise FileNotFoundError(f"Workbook not found at {file_path}")

    ext = file_path.suffix.lower()
    if ext == ".xls":
        return _parse_xls_workbook(file_path)

    if ext != ".xlsx":
        raise ValueError(f"Unsupported file extension '{ext}'. Only .xlsx and .xls are supported.")

    wb = openpyxl.load_workbook(str(file_path), data_only=True)
    all_rows: List[ParsedRow] = []
    sheets_summary: List[SheetSummary] = []
    total_discovered_rows = 0
    total_imported_rows = 0
    total_failed_rows = 0

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        sheet_rows, eligible, imported, failed = _parse_openpyxl_sheet(sheet, sheet_name)
        all_rows.extend(sheet_rows)
        sheets_summary.append(SheetSummary(sheet_name, eligible, imported, failed))
        total_discovered_rows += eligible
        total_imported_rows += imported
        total_failed_rows += failed

    return NcrpParseResult(
        sheet_count=len(sheets_summary),
        total_rows=total_discovered_rows,
        imported_rows=total_imported_rows,
        failed_rows=total_failed_rows,
        rows=all_rows,
        sheets_summary=sheets_summary,
    )
