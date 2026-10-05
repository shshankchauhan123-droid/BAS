import os
import json
import urllib.request
import urllib.error
import pandas as pd

from app.processing.mapping.qwen_header_prompt import build_qwen_header_prompt
from app.processing.mapping.header_normalize import normalize_header



def map_headers_with_qwen(normalized_headers: list[str], raw_excel_path: str, header_row_index, is_excel=False):
    """
    Sends normalized headers to Qwen, validates the response, 
    and updates the raw Excel headers without modifying data.
    """
    
    ALLOWED_FIELDS = {
        "transaction_date",
        "description",
        "cheque_number",
        "debit",
        "credit",
        "balance",
        "mode"
    }
    
    # Account metadata is extracted separately in Stage 1 and stored in the file record.
    # Qwen should only map transaction-level fields.

    qwen_model = os.getenv("QWEN_MODEL", "qwen3.5:0.8b")
    ollama_url = os.getenv("OLLAMA_BASE_URL", "http://localhost:11434")
    
    try:
        confidence_threshold = float(os.getenv("QWEN_HEADER_CONFIDENCE_THRESHOLD", "0.80"))
    except ValueError:
        confidence_threshold = 0.80

    print("============================================================")
    print("STAGE 5 - QWEN HEADER MAPPING")
    print("============================================================")
    print()
    print(f"QWEN MODEL: {qwen_model}")
    print(f"OLLAMA URL: {ollama_url}")
    print()
    print("INPUT HEADERS:")
    print(normalized_headers)
    print()
    
    # ---------------------------------------------------------
    # DETERMINISTIC MAPPING
    # ---------------------------------------------------------
    DETERMINISTIC_ALIASES = {
        "transaction_date": ["date", "txn date", "value date", "transaction date", "value dt", "txn dt"],
        "description": ["description", "particulars", "narration", "remarks", "transaction details"],
        "debit": ["debit", "withdrawal", "withdrawals", "withdrawal amt.", "withdrawal amt", "dr", "dr amount"],
        "credit": ["credit", "deposit", "deposits", "deposit amt.", "deposit amt", "cr", "cr amount"],
        "balance": ["balance", "closing balance", "available balance", "account balance"],
        "cheque_number": ["chq", "cheque", "ref no", "instrument no", "cheque number", "reference number", "chq./ref.no."],
        "mode": ["mode", "type", "transaction mode", "txn type"]
    }
    
    valid_map_dict = {}
    used_target_fields = set()
    unmapped_headers = []
    
    print("DETERMINISTIC MAPPING PHASE:")
    for header in normalized_headers:
        mapped = False
        for target, aliases in DETERMINISTIC_ALIASES.items():
            if header in aliases:
                if target not in used_target_fields:
                    valid_map_dict[header] = target
                    used_target_fields.add(target)
                    print(f"{header.ljust(15)} -> {target.ljust(20)} [Deterministic]")
                    mapped = True
                    break
        if not mapped:
            # Fallback: check if any alias is a substring of the header
            for target, aliases in DETERMINISTIC_ALIASES.items():
                if any(alias in header for alias in aliases):
                    if target not in used_target_fields:
                        valid_map_dict[header] = target
                        used_target_fields.add(target)
                        print(f"{header.ljust(15)} -> {target.ljust(20)} [Deterministic Partial]")
                        mapped = True
                        break
        
        if not mapped:
            unmapped_headers.append(header)
            
    print()
    
    if unmapped_headers:
        print("Sending unmapped headers to Qwen...\n")
        print("UNMAPPED HEADERS:")
        print(unmapped_headers)
        print()
    
        prompt = build_qwen_header_prompt(unmapped_headers, list(ALLOWED_FIELDS - used_target_fields))
    
        try:
            req = urllib.request.Request(
                f"{ollama_url}/api/generate",
                data=json.dumps({
                    "model": qwen_model,
                    "prompt": prompt,
                    "stream": False,
                    "think": False,
                    "format": "json",
                    "options": {
                        "temperature": 0,
                        "num_predict": 2048
                    }
                }).encode('utf-8'),
                headers={'Content-Type': 'application/json'},
                method='POST'
            )
            
            with urllib.request.urlopen(req, timeout=500) as response:
                raw_http_response = response.read().decode('utf-8')
                
        except Exception as e:
            raise RuntimeError(f"Failed to call Ollama API: {str(e)}")

    if unmapped_headers:
        print("QWEN RESPONSE RECEIVED\n")
        
        try:
            data = json.loads(raw_http_response)
        except json.JSONDecodeError:
            raise ValueError("Ollama HTTP response is not valid JSON.")
            
        print("=" * 60)
        print("FULL OLLAMA API RESPONSE")
        print(raw_http_response)
        print("=" * 60)
        print()
        
        response_text = data.get("response", "")
        
        if not response_text.strip():
            raise ValueError(
                f"Qwen returned an empty response. "
                f"done={data.get('done')}, "
                f"done_reason={data.get('done_reason')}, "
                f"eval_count={data.get('eval_count')}, "
                f"prompt_eval_count={data.get('prompt_eval_count')}"
            )
    
        print("QWEN RESPONSE RECEIVED\n")
        
        print("=" * 60)
        print("RAW QWEN RESPONSE")
        print(response_text)
        print("=" * 60)
        print()
        
        try:
            qwen_json = json.loads(response_text)
        except json.JSONDecodeError:
            import re
            match = re.search(r"```(?:json)?\s*(.*?)\s*```", response_text, re.DOTALL | re.IGNORECASE)
            if match:
                try:
                    qwen_json = json.loads(match.group(1).strip())
                except json.JSONDecodeError:
                    raise ValueError("Qwen response does not contain valid JSON inside markdown.")
            else:
                match = re.search(r"(\{.*\})", response_text, re.DOTALL)
                if match:
                    try:
                        qwen_json = json.loads(match.group(1).strip())
                    except json.JSONDecodeError:
                        raise ValueError("Qwen response is not valid JSON.")
                else:
                    raise ValueError("Qwen response is not valid JSON.")
    
        if not isinstance(qwen_json, dict) or "mappings" not in qwen_json:
            raise ValueError("Qwen response missing 'mappings' array.")
            
        mappings = qwen_json.get("mappings")
        if not isinstance(mappings, list):
            raise ValueError("'mappings' must be an array.")
            
        print("QWEN MAPPING:\n")
        
        # Sort mappings by confidence descending
        mappings.sort(
            key=lambda x: float(x.get("confidence", 0)) 
            if str(x.get("confidence", 0)).replace('.', '', 1).isdigit() 
            else 0.0, 
            reverse=True
        )
    
        for mapping in mappings:
            if not isinstance(mapping, dict):
                continue
                
            source_header = mapping.get("source_header")
            target_field = mapping.get("target_field")
            confidence = mapping.get("confidence")
            
            if source_header and target_field:
                try:
                    conf_val = float(confidence)
                    conf_str = f"[{conf_val:.2f}]"
                except (TypeError, ValueError):
                    conf_str = "[N/A]"
                print(f"{str(source_header).ljust(15)} -> {str(target_field).ljust(20)} {conf_str}")
                
            if not source_header or not target_field:
                continue
                
            if source_header not in unmapped_headers:
                continue
                
            if target_field not in ALLOWED_FIELDS:
                continue
                
            try:
                confidence = float(confidence)
            except (TypeError, ValueError):
                confidence = 0.0
                
            if confidence < confidence_threshold:
                print(f"Skipping mapping for '{source_header}' due to low confidence ({confidence} < {confidence_threshold})")
                continue
                
            if target_field in used_target_fields:
                print(f"Skipping mapping for '{source_header}' -> '{target_field}' because target field is already mapped by a higher confidence source.")
                continue
                
            valid_map_dict[source_header] = target_field
            used_target_fields.add(target_field)

    print("\nValidating Qwen mapping...\n")
    
    if not valid_map_dict:
        print("Warning: No high-confidence valid mappings were found.")
    else:
        print("Mapping validation successful.\n")
        
    
    absolute_excel_path = os.path.abspath(raw_excel_path)
    
    print("EXCEL FILE BEING UPDATED:")
    print(absolute_excel_path)
    print()
    
    if not os.path.exists(absolute_excel_path):
        raise FileNotFoundError(f"Raw Excel file not found: {absolute_excel_path}")
        
    try:
        import openpyxl
        import shutil
        
        workbook = openpyxl.load_workbook(absolute_excel_path)
        sheet = workbook.active
        
        original_max_row = sheet.max_row
        original_max_column = sheet.max_column
        
        excel_header_row = 1 if header_row_index == "columns" else int(header_row_index) + 2
        
        print("EXCEL HEADER ROW:")
        print(excel_header_row)
        print()
        
        # Read all rows to verify later
        all_original_rows = []
        for r in range(1, original_max_row + 1):
            all_original_rows.append([sheet.cell(row=r, column=c).value for c in range(1, original_max_column + 1)])
            
        original_headers = all_original_rows[excel_header_row - 1]
        
        print("EXCEL HEADERS BEFORE UPDATE:")
        print(original_headers)
        print()
        
        for c in range(1, original_max_column + 1):
            cell = sheet.cell(row=excel_header_row, column=c)
            val = cell.value
            if val:
                norm_val = normalize_header(str(val))
                if norm_val in valid_map_dict:
                    cell.value = valid_map_dict[norm_val]
        
        # Atomic save
        temp_excel_path = absolute_excel_path + ".tmp.xlsx"
        workbook.save(temp_excel_path)
        workbook.close()
        
        # Reload and verify from temp file
        wb_check = openpyxl.load_workbook(temp_excel_path)
        sh_check = wb_check.active
        
        if sh_check.max_row != original_max_row:
            wb_check.close()
            os.remove(temp_excel_path)
            raise ValueError(f"Excel row count changed from {original_max_row} to {sh_check.max_row}")
            
        if sh_check.max_column != original_max_column:
            wb_check.close()
            os.remove(temp_excel_path)
            raise ValueError(f"Excel column count changed from {original_max_column} to {sh_check.max_column}")
            
        new_headers = [sh_check.cell(row=excel_header_row, column=c).value for c in range(1, original_max_column + 1)]
        
        print("EXCEL HEADERS AFTER UPDATE:")
        print(new_headers)
        print()
        
        # Verify mapped headers
        headers_changed = False
        for i, val in enumerate(original_headers):
            if val:
                norm_val = normalize_header(str(val))
                if norm_val in valid_map_dict:
                    expected_target = valid_map_dict[norm_val]
                    if new_headers[i] != expected_target:
                        wb_check.close()
                        os.remove(temp_excel_path)
                        raise ValueError(f"Header '{val}' was not correctly updated to '{expected_target}'")
                    if original_headers[i] != expected_target:
                        headers_changed = True
        
        if not headers_changed:
            wb_check.close()
            os.remove(temp_excel_path)
            raise ValueError("No headers were actually changed in the Excel file.")
            
        # Verify all other rows
        for r in range(1, original_max_row + 1):
            if r == excel_header_row:
                continue
            current_vals = [sh_check.cell(row=r, column=c).value for c in range(1, original_max_column + 1)]
            if current_vals != all_original_rows[r - 1]:
                wb_check.close()
                os.remove(temp_excel_path)
                raise ValueError(f"Transaction data at row {r} was modified during header update!")
        
        wb_check.close()
        
        # Replace original
        os.replace(temp_excel_path, absolute_excel_path)
        
        print("EXCEL FILE UPDATED SUCCESSFULLY:")
        print(absolute_excel_path)
        print()
        
    except Exception as e:
        raise RuntimeError(f"Failed to update Excel file: {str(e)}")

    print("STAGE 5 SUCCESS\n")
    print("STATUS: PROCESSING")
    print("STAGE: qwen_header_mapping")
    print("PROGRESS: 50%\n")
